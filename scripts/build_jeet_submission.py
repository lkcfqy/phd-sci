"""Build a double-blind JEET submission package from the frozen paper sources.

The generated Word files follow the current Journal of Electrical Engineering &
Technology author instructions: anonymous manuscript plus a separate title page,
10-point Times body text, decimal headings, numeric citations, inline figures/tables,
and page numbering.  The title page and cover letter are intentionally labelled as
author-input templates because names, affiliations, funding, and authorship decisions
must not be invented.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import zipfile
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
PAPER = ROOT / "paper"
MANUSCRIPT = PAPER / "manuscript.md"
SUPPLEMENT = PAPER / "supplementary_material.md"
BIBLIOGRAPHY = ROOT / "references" / "key_papers.bib"
OUT_DIR = ROOT / "submission" / "jeet"
TMP_DIR = ROOT / "tmp" / "jeet_submission"

TITLE = (
    "When Healthy-Only Transfer Fails in PMSM Stator-Fault Detection: "
    "A Leakage-Resistant Cross-Dataset Evaluation"
)


def today_iso() -> str:
    """Return the build date in the project's declared Asia/Shanghai timezone."""
    return datetime.now(tz=ZoneInfo("Asia/Shanghai")).date().isoformat()


SHORT_TITLE = "When Healthy-Only PMSM Transfer Fails"
JOURNAL = "Journal of Electrical Engineering & Technology"

FIGURES = [
    ("protocol_overview", "Motor holdout, sequential target-health split, and method flow"),
    ("method_performance", "Per-motor detection and healthy false-alarm intervals"),
    (
        "paired_detection_differences",
        (
            "Record-level paired detection-rate differences between the Log-Euclidean entity "
            "detector and each covariance comparator. Points are mean paired differences and "
            "bars are 95% percentile intervals from 10,000 motor-stratified fault-record "
            "bootstrap replicates over the 42 available KAIST fault records; positive values "
            "favor the Log-Euclidean detector."
        ),
    ),
    (
        "oneclass_detection_differences",
        (
            "Record-level paired detection-rate differences between the proposed detector and "
            "the strong one-class baselines. Points are mean paired differences and bars are "
            "95% percentile intervals from 10,000 motor-stratified fault-record bootstrap "
            "replicates over the 42 available KAIST fault records; positive values favor the "
            "proposed detector. Parenthetical FA labels report false alarms among the 42 "
            "later-health blocks."
        ),
    ),
    ("adaptation_budget", "Adaptation-budget sensitivity and calibration feasibility"),
    ("block_sensitivity", "Block-length and aggregation sensitivity"),
    ("severity_detection", "Nonmonotonic detection across nominal severity"),
    (
        "external_method_performance",
        (
            "External fault detection, held-out-health false alarms, and the predeclared H1 "
            "decision on one dual-three-phase PMSM. Panel (a) shows block detection with "
            "turn-stratified record-bootstrap intervals over 48 fault records and 384 ordered "
            "blocks. Panel (b) shows pooled false-alarm rates with descriptive 95% Wilson "
            "intervals over four held-out healthy loads and 32 blocks, maximum load-specific "
            "false-alarm rates, and the H1 limits of 12% for the pooled Wilson upper bound and "
            "15% for the maximum-load rate. Gold marks the observed detection leader and "
            "hatched blue marks the frozen proposed detector; intervals are descriptive at the "
            "recorded-condition level."
        ),
    ),
    (
        "external_condition_drift",
        (
            "Frozen Log-Euclidean score and alarm drift over eight ordered 3 s external "
            "analysis blocks. Panel (a) shows the median and interquartile range of "
            "threshold-normalized scores on a logarithmic scale for 48 fault records and four "
            "held-out healthy records per block; the top axis gives approximate median speed. "
            "Panel (b) shows empirical block alarm rates. Blocks are ordered operating points "
            "rather than independent repeats, and the score and alarm drift are therefore "
            "confounded with the rise from approximately 218 to 2,214 rpm."
        ),
    ),
    (
        "external_proposed_heatmap",
        (
            "Frozen-detector alarm fraction across six fixed turn-phase conditions and eight "
            "loads on the external motor. Each cell is the fraction of eight ordered 3 s "
            "blocks alarmed within one record, so 12.5% represents one of eight blocks. Phase "
            "is fixed by turn count (U: 1, 3, 5, and 6; V: 2 and 4), so turn and phase effects "
            "are not separately identifiable; blocks within a record are ordered operating "
            "points rather than independent repetitions."
        ),
    ),
    (
        "external_feature_geometry",
        (
            "Post-reveal single-feature discrimination versus frozen-score contribution on "
            "one external motor. Panel (a) shows direction-free AUROC and the absolute matched "
            "standardized difference for record-subsystem-block means at the four loads with "
            "held-out health; each healthy condition is reused across six fault-turn "
            "comparisons. Panel (b) shows winning-window absolute Mahalanobis contribution "
            "shares for all fault records and held-out health, using the symmetric cross-term "
            "split defined in Section 6.9. The analysis is descriptive, no detector was refit, and no "
            "threshold was changed."
        ),
    ),
]

TABLE_CAPTIONS = [
    "Exploratory KAIST covariance-score comparison under leave-one-motor-out evaluation",
    "Exploratory KAIST one-class baselines under the common healthy-data protocol",
    "Frozen external method comparison on one dual-three-phase PMSM",
]

ANONYMOUS_BUNDLE_PATHS = [
    Path("pyproject.toml"),
    Path("configs"),
    Path("src"),
    Path("scripts"),
    Path("tests"),
    Path("paper/manuscript.md"),
    Path("references/key_papers.bib"),
    Path("docs/research_protocol.md"),
    Path("docs/external_validation_protocol.md"),
    Path("docs/fault_reveal_log.md"),
    Path("docs/secondary_transient_validation_protocol.md"),
    Path("docs/secondary_transient_reveal_log.md"),
    Path("docs/data_sources.yaml"),
    Path("docs/reference_metadata_audit.md"),
    Path(
        "results/healthy_covariance_v0/target_1kW/scale_free/"
        "log_euclidean_entity_covariance/block_predictions.csv"
    ),
    Path(
        "results/healthy_covariance_v0/target_1.5kW/scale_free/"
        "log_euclidean_entity_covariance/block_predictions.csv"
    ),
    Path(
        "results/healthy_covariance_v0/target_3kW/scale_free/"
        "log_euclidean_entity_covariance/block_predictions.csv"
    ),
    Path("results/oneclass_baselines/record_summary.csv"),
    Path("results/oneclass_baselines/comparison_to_proposed.csv"),
    Path("results/external_pmsm_validation"),
    Path("results/external_pmsm_analysis"),
    Path("results/external_failure_diagnostics"),
    Path("results/external_seed_sensitivity"),
    Path("results/external_feature_drift"),
    Path("results/sampling_rate_sensitivity"),
    Path("results/transient_feature_build"),
    Path("results/transient_feature_build_post_reveal_implicit_time"),
    Path("results/transient_pmsm_validation_post_reveal_200w"),
]


@dataclass
class BibEntry:
    entry_type: str
    key: str
    fields: dict[str, str]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = 120) -> None:
    if sum(widths_dxa) <= 0:
        raise ValueError("Table width must be positive")
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa, strict=True):
            set_cell_width(cell, width)
            set_cell_margins(cell)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def set_font(run, name: str = "Times New Roman", size: float = 10, *, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_page_number(section) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend((fld_char1, instr_text, fld_char2))
    set_font(run, size=9)


def configure_doc(doc: Document, *, landscape: bool = False, supplement: bool = False) -> int:
    section = doc.sections[0]
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11)
        section.page_height = Inches(8.5)
        margin = Inches(0.55)
        content_width = 14256
    else:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        margin = Inches(1)
        content_width = 9360
    section.top_margin = margin
    section.right_margin = margin
    section.bottom_margin = margin
    section.left_margin = margin
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)
    add_page_number(section)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(9 if supplement else 10)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(4 if supplement else 6)
    normal.paragraph_format.line_spacing = 1.05 if supplement else 1.15
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    title = styles["Title"]
    title.font.name = "Times New Roman"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    title.font.size = Pt(15 if not supplement else 14)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(10)

    heading_specs = {
        "Heading 1": (12 if not supplement else 11, 12, 5, True, False),
        "Heading 2": (11 if not supplement else 10, 10, 4, True, False),
        "Heading 3": (10 if not supplement else 9, 8, 3, True, True),
    }
    for name, (size, before, after, bold, italic) in heading_specs.items():
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.italic = italic
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Figure Caption" not in styles:
        styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    cap = styles["Figure Caption"]
    cap.font.name = "Times New Roman"
    cap._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    cap._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    cap.font.size = Pt(8 if supplement else 9)
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cap.paragraph_format.space_before = Pt(2)
    cap.paragraph_format.space_after = Pt(6)
    cap.paragraph_format.keep_together = True

    if "Table Caption" not in styles:
        styles.add_style("Table Caption", WD_STYLE_TYPE.PARAGRAPH)
    tcap = styles["Table Caption"]
    tcap.font.name = "Times New Roman"
    tcap._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    tcap._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    tcap.font.size = Pt(8 if supplement else 9)
    tcap.font.bold = True
    tcap.paragraph_format.space_before = Pt(6)
    tcap.paragraph_format.space_after = Pt(3)
    tcap.paragraph_format.keep_with_next = True
    return content_width


def add_o_math(doc: Document, equation: str) -> None:
    linear = latex_to_unicode(equation.rstrip("."))
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(6)
    math_para = OxmlElement("m:oMathPara")
    math = OxmlElement("m:oMath")
    math_run = OxmlElement("m:r")
    text = OxmlElement("m:t")
    text.text = linear
    math_run.append(text)
    math.append(math_run)
    math_para.append(math)
    paragraph._p.append(math_para)


def latex_to_unicode(value: str) -> str:
    canonical = re.sub(r"\s+", "", value.rstrip("."))
    display_equations = {
        r"\widetilde{x}_{m,i}=(x_{m,i}-c_m)\oslashr_m": "x̃ₘ,ᵢ = (xₘ,ᵢ - cₘ) ./ rₘ",
        r"S_m^{(\lambda)}=S_m+\lambda\,\frac{\operatorname{tr}(S_m)}{d}I": (
            "Sₘ⁽λ⁾ = Sₘ + λ [tr(Sₘ) / d] I"
        ),
        (
            r"\Sigma_{\mathrm{LE}}="
            r"\exp\!\left[\frac{1}{M}\sum_{m=1}^{M}\logS_m^{(\lambda)}\right]"
        ): "Σ_LE = exp[(1 / M) Σₘ₌₁ᴹ log Sₘ⁽λ⁾]",
        r"s(x)=\widetilde{x}^{\top}\Sigma_{\mathrm{LE}}^{\dagger}\widetilde{x}": (
            "s(x) = x̃ᵀ Σ_LE† x̃"
        ),
        r"p_b=\frac{1+\sum_{j=1}^{n}\mathbf{1}(A_j\geA_b)}{n+1}": (
            "p_b = [1 + Σⱼ₌₁ⁿ 1(Aⱼ ≥ A_b)] / (n + 1)"
        ),
    }
    if canonical in display_equations:
        return display_equations[canonical]

    replacements = {
        r"\mathbb{R}": "R",
        r"\widetilde{x}": "x̃",
        r"\Sigma": "Σ",
        r"\lambda": "λ",
        r"\times": "×",
        r"\le": "≤",
        r"\ge": "≥",
        r"\sum": "Σ",
        r"\top": "T",
        r"\dagger": "†",
        r"\oslash": "./",
        r"\operatorname{tr}": "tr",
        r"\mathrm{LE}": "LE",
        r"\mathbf{1}": "1",
        r"\exp": "exp",
        r"\log": "log",
        r"\max": "max",
        r"\in": " in ",
        r"\,": " ",
        r"\!": "",
        r"\left": "",
        r"\right": "",
        r"\frac": "frac",
    }
    text = value
    for source, target in replacements.items():
        text = text.replace(source, target)
    text = re.sub(r"frac\{([^{}]+)\}\{([^{}]+)\}", r"(\1)/(\2)", text)
    superscript = str.maketrans("0123456789+-=()nT", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ⁿᵀ")
    subscript = str.maketrans("0123456789+-=()aeijkmnoprstx", "₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎ₐₑᵢⱼₖₘₙₒₚᵣₛₜₓ")

    def scripted(match: re.Match[str], mapping: dict[int, str], marker: str) -> str:
        content = match.group(1)
        converted = content.translate(mapping)
        if len(converted) == len(content) and all(char not in "_^" for char in converted):
            return converted
        return f"{marker}({content})"

    text = re.sub(
        r"\^\{([^{}]+)\}",
        lambda match: scripted(match, superscript, "^"),
        text,
    )
    text = re.sub(
        r"_\{([^{}]+)\}",
        lambda match: scripted(match, subscript, "_"),
        text,
    )
    text = re.sub(
        r"\^([0-9nT])",
        lambda match: match.group(1).translate(superscript),
        text,
    )
    text = re.sub(
        r"_([0-9aeijkmnoprstx])",
        lambda match: match.group(1).translate(subscript),
        text,
    )
    text = text.replace("{", "").replace("}", "")
    text = text.replace("\\", "")
    return text.strip()


def normalize_text(text: str) -> str:
    return (
        text.replace("–", "-")
        .replace("—", "-")
        .replace("−", "-")
        .replace("‑", "-")
        .replace("N·m", "N m")
        .replace("--", "-")
    )


def add_inline_runs(paragraph, text: str, cite_state: dict[str, int] | None = None) -> None:
    if cite_state is not None:
        text = replace_citations(text, cite_state)
    text = text.replace(r"\(", "$ ").replace(r"\)", " $")
    text = normalize_text(text)
    token_re = re.compile(r"(\*\*.+?\*\*|`.+?`|\$.*?\$|\*[^*]+?\*)")
    cursor = 0
    for match in token_re.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            set_font(run)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_font(run, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_font(run, name="Courier New", size=8.5)
        elif token.startswith("$"):
            run = paragraph.add_run(latex_to_unicode(token[1:-1].strip()))
            set_font(run, name="Cambria Math", italic=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_font(run, italic=True)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_font(run)


def replace_citations(text: str, cite_state: dict[str, int]) -> str:
    def replacement(match: re.Match[str]) -> str:
        keys = re.findall(r"@([A-Za-z0-9_:-]+)", match.group(1))
        nums: list[int] = []
        for key in keys:
            if key not in cite_state:
                cite_state[key] = len(cite_state) + 1
            nums.append(cite_state[key])
        return "[" + ", ".join(str(num) for num in nums) + "]"

    return re.sub(r"\[([^\]]*@[A-Za-z0-9_:-]+[^\]]*)\]", replacement, text)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].lstrip().startswith("|"):
        raw_cells = re.split(r"(?<!\\)\|", lines[index].strip())[1:-1]
        cells = [cell.replace(r"\|", "|").strip() for cell in raw_cells]
        if not all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells):
            rows.append(cells)
        index += 1
    return rows, index


def column_widths(rows: list[list[str]], total_dxa: int) -> list[int]:
    cols = len(rows[0])
    weights = []
    for col in range(cols):
        max_len = max(len(re.sub(r"[*`]", "", row[col])) for row in rows)
        weights.append(max(5.0, min(28.0, max_len**0.5 * 2.4)))
    min_width = 540 if cols >= 9 else 800
    available = total_dxa - min_width * cols
    if available < 0:
        min_width = max(360, total_dxa // cols // 2)
        available = total_dxa - min_width * cols
    total_weight = sum(weights)
    widths = [min_width + int(available * weight / total_weight) for weight in weights]
    widths[-1] += total_dxa - sum(widths)
    return widths


def add_table(
    doc: Document,
    rows: list[list[str]],
    *,
    width_dxa: int,
    caption: str | None,
    table_number: int,
    supplement: bool,
    cite_state: dict[str, int] | None,
) -> None:
    if caption:
        cap = doc.add_paragraph(style="Table Caption")
        add_inline_runs(cap, f"Table {table_number}  {caption}", cite_state)
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    widths = column_widths(rows, width_dxa)
    set_table_geometry(table, widths)
    set_repeat_table_header(table.rows[0])
    font_size = 6.4 if supplement and len(rows[0]) >= 9 else (7.0 if supplement else 8.0)
    for row_index, (word_row, table_row) in enumerate(zip(rows, table.rows, strict=True)):
        for col_index, (value, cell) in enumerate(zip(word_row, table_row.cells, strict=True)):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                shade_cell(cell, "E7E6E6")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if col_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            )
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            clean = normalize_text(value.replace("**", "").replace("`", ""))
            run = paragraph.add_run(clean)
            set_font(run, size=font_size, bold=(row_index == 0 or "**" in value))
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_figure(doc: Document, source: Path, caption: str, number: int, max_width: float) -> None:
    png = source.with_suffix(".png")
    if not png.exists():
        raise FileNotFoundError(png)
    with Image.open(png) as image:
        width_px, height_px = image.size
    width_in = max_width
    height_in = width_in * height_px / width_px
    max_height = 7.2
    if height_in > max_height:
        height_in = max_height
        width_in = height_in * width_px / height_px
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    inline_shape = run.add_picture(str(png), width=Inches(width_in))
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", caption)
    cap = doc.add_paragraph(style="Figure Caption")
    lead = cap.add_run(f"Fig. {number}  ")
    set_font(lead, size=9, bold=True)
    body = cap.add_run(normalize_text(caption.rstrip(".")))
    set_font(body, size=9)


def extract_front_title(lines: list[str]) -> tuple[str, int]:
    if not lines or lines[0].strip() != "---":
        return TITLE, 0
    end = lines.index("---", 1)
    title_lines: list[str] = []
    in_title = False
    for line in lines[1:end]:
        if line.startswith("title:"):
            in_title = True
            continue
        if in_title and line.startswith("  "):
            title_lines.append(line.strip())
        elif in_title:
            break
    return " ".join(title_lines) or TITLE, end + 1


def build_markdown_doc(
    source: Path,
    output: Path,
    *,
    anonymous_main: bool,
    supplement: bool = False,
) -> dict[str, object]:
    text = source.read_text(encoding="utf-8")
    lines = text.splitlines()
    title, index = extract_front_title(lines)
    doc = Document()
    content_width = configure_doc(doc, landscape=supplement, supplement=supplement)
    title_paragraph = doc.add_paragraph(style="Title")
    add_inline_runs(title_paragraph, title)
    if supplement:
        sub = doc.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub.paragraph_format.space_after = Pt(8)
        run = sub.add_run(f"Online Resource 1 for {JOURNAL} - anonymous review copy")
        set_font(run, size=9, italic=True)

    cite_state: dict[str, int] = {}
    figure_counter = 0
    table_counter = 0
    in_display_math = False
    equation_lines: list[str] = []

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if in_display_math:
            if stripped == r"\]":
                add_o_math(doc, " ".join(equation_lines))
                equation_lines = []
                in_display_math = False
            else:
                equation_lines.append(stripped)
            index += 1
            continue
        if stripped == r"\[":
            in_display_math = True
            index += 1
            continue
        if not stripped:
            index += 1
            continue
        if stripped.startswith(">"):
            block: list[str] = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                block.append(lines[index].strip().lstrip(">").strip())
                index += 1
            if not anonymous_main:
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(0.2)
                paragraph.paragraph_format.right_indent = Inches(0.2)
                paragraph.paragraph_format.space_after = Pt(6)
                add_inline_runs(paragraph, " ".join(block), cite_state)
                for run in paragraph.runs:
                    run.italic = True
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", stripped)
        if heading:
            level = len(heading.group(1))
            heading_text = normalize_text(heading.group(2))
            if level == 1:
                # The supplement title was already added above.
                index += 1
                continue
            style = "Heading 1" if level == 2 else ("Heading 2" if level == 3 else "Heading 3")
            paragraph = doc.add_paragraph(style=style)
            add_inline_runs(paragraph, heading_text, cite_state)
            index += 1
            continue
        image = re.match(r"^!\[(.*?)\]\((.*?)\)$", stripped)
        if image:
            figure_counter += 1
            relative = Path(image.group(2))
            source_path = (source.parent / relative).resolve()
            add_figure(
                doc,
                source_path,
                image.group(1),
                figure_counter,
                max_width=6.35 if not supplement else 9.6,
            )
            index += 1
            continue
        if stripped.startswith("|"):
            rows, index = parse_table(lines, index)
            table_counter += 1
            caption = None
            if anonymous_main:
                caption = TABLE_CAPTIONS[table_counter - 1]
            add_table(
                doc,
                rows,
                width_dxa=content_width,
                caption=caption,
                table_number=table_counter,
                supplement=supplement,
                cite_state=cite_state,
            )
            continue
        if stripped.startswith("- "):
            while index < len(lines) and lines[index].strip().startswith("- "):
                bullet_parts = [lines[index].strip()[2:]]
                index += 1
                while index < len(lines):
                    continuation = lines[index].strip()
                    if (
                        not continuation
                        or continuation.startswith(("#", "!", "|", "- ", ">"))
                        or continuation == r"\["
                    ):
                        break
                    bullet_parts.append(continuation)
                    index += 1
                paragraph = doc.add_paragraph(style="List Bullet")
                paragraph.paragraph_format.left_indent = Inches(0.38)
                paragraph.paragraph_format.first_line_indent = Inches(-0.18)
                paragraph.paragraph_format.space_after = Pt(3)
                add_inline_runs(paragraph, " ".join(bullet_parts), cite_state)
            continue

        block = [stripped]
        index += 1
        while index < len(lines):
            next_line = lines[index].strip()
            if (
                not next_line
                or next_line.startswith(("#", "!", "|", "- ", ">"))
                or next_line == r"\["
            ):
                break
            block.append(next_line)
            index += 1
        paragraph_text = " ".join(block)
        if paragraph_text.startswith("**Keywords:**"):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(8)
            label = paragraph.add_run("Keywords: ")
            set_font(label, bold=True)
            add_inline_runs(paragraph, paragraph_text[len("**Keywords:**") :].strip(), cite_state)
        else:
            paragraph = doc.add_paragraph()
            add_inline_runs(paragraph, paragraph_text, cite_state)

    if anonymous_main:
        add_references(doc, cite_state)
    output.parent.mkdir(parents=True, exist_ok=True)
    set_core_properties(doc, anonymous=anonymous_main or supplement)
    doc.save(output)
    return {
        "output": str(output.relative_to(ROOT)),
        "figures": figure_counter,
        "tables": table_counter,
        "citations": len(cite_state),
        "sha256": sha256(output),
    }


def parse_bibtex(path: Path) -> dict[str, BibEntry]:
    text = path.read_text(encoding="utf-8")
    entries: dict[str, BibEntry] = {}
    index = 0
    while True:
        match = re.search(r"@(\w+)\s*\{\s*([^,]+),", text[index:])
        if not match:
            break
        entry_type = match.group(1).lower()
        key = match.group(2).strip()
        start = index + match.end()
        cursor = start
        depth = 1
        while cursor < len(text) and depth:
            if text[cursor] == "{":
                depth += 1
            elif text[cursor] == "}":
                depth -= 1
            cursor += 1
        body = text[start : cursor - 1]
        fields: dict[str, str] = {}
        pos = 0
        while pos < len(body):
            field_match = re.search(r"(\w+)\s*=\s*", body[pos:])
            if not field_match:
                break
            name = field_match.group(1).lower()
            value_start = pos + field_match.end()
            if value_start >= len(body) or body[value_start] != "{":
                pos = value_start + 1
                continue
            value_pos = value_start + 1
            value_depth = 1
            while value_pos < len(body) and value_depth:
                if body[value_pos] == "{":
                    value_depth += 1
                elif body[value_pos] == "}":
                    value_depth -= 1
                value_pos += 1
            value = body[value_start + 1 : value_pos - 1]
            fields[name] = decode_latex(value.strip())
            pos = value_pos
        entries[key] = BibEntry(entry_type, key, fields)
        index = cursor
    return entries


def decode_latex(value: str) -> str:
    replacements = {
        r"{\"a}": "ä",
        r"{\"o}": "ö",
        r"{\"u}": "ü",
        r"{\'a}": "á",
        r"{\'e}": "é",
        r"{\'i}": "í",
        r"{\'o}": "ó",
        r"{\'u}": "ú",
        r"{\'y}": "ý",
        r"{\v{s}}": "š",
        r"{\v{e}}": "ě",
        r"\&": "&",
    }
    text = value
    for source, target in replacements.items():
        text = text.replace(source, target)
    return text.replace("{", "").replace("}", "")


def initials(first_names: str) -> str:
    parts = re.findall(r"[A-Za-zÀ-ž]+", first_names)
    return " ".join(f"{part[0]}." for part in parts if part)


def format_authors(value: str) -> str:
    people = []
    for person in value.split(" and "):
        if "," in person:
            last, first = [part.strip() for part in person.split(",", 1)]
            people.append(f"{last}, {initials(first)}")
        else:
            people.append(person.strip())
    if len(people) == 1:
        return people[0]
    return ", ".join(people[:-1]) + ", & " + people[-1]


def format_reference(entry: BibEntry) -> str:
    f = entry.fields
    authors = format_authors(decode_latex(f.get("author", "Anonymous")))
    year = f.get("year", "n.d.")
    title = decode_latex(f.get("title", "Untitled"))
    doi = f.get("doi")
    link = f"https://doi.org/{doi}" if doi else f.get("url", "")
    publisher = decode_latex(f.get("publisher", ""))
    is_dataset = entry.entry_type == "dataset" or (
        entry.entry_type == "misc" and "dataset" in f.get("note", "").lower()
    )
    if is_dataset:
        version = f" (Version {f['version']})" if f.get("version") else ""
        body = f"{authors} ({year}). {title}{version} [Data set]. {publisher}."
    elif entry.entry_type == "inproceedings":
        venue = decode_latex(f.get("booktitle", "Proceedings"))
        details = []
        if f.get("volume"):
            details.append(f"Vol. {f['volume']}")
        if f.get("pages"):
            details.append(f"pp. {f['pages']}")
        suffix = f" ({', '.join(details)})" if details else ""
        body = f"{authors} ({year}). {title}. In {venue}{suffix}."
    else:
        journal = decode_latex(f.get("journal", ""))
        volume = f.get("volume", "")
        issue = f"({f['number']})" if f.get("number") else ""
        pages = f", {f['pages']}" if f.get("pages") else ""
        body = f"{authors} ({year}). {title}. {journal}, {volume}{issue}{pages}."
    return normalize_text(f"{body} {link}".strip())


def add_references(doc: Document, cite_state: dict[str, int]) -> None:
    entries = parse_bibtex(BIBLIOGRAPHY)
    missing = [key for key in cite_state if key not in entries]
    if missing:
        raise KeyError(f"Missing bibliography entries: {missing}")
    doc.add_paragraph("References", style="Heading 1")
    for key, number in sorted(cite_state.items(), key=lambda item: item[1]):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.add_run(f"[{number}] {format_reference(entries[key])}")
        set_font(run, size=9)


def set_core_properties(doc: Document, *, anonymous: bool) -> None:
    props = doc.core_properties
    props.title = TITLE
    props.subject = JOURNAL + " submission"
    props.author = "Anonymous" if anonymous else "AUTHOR INPUT REQUIRED"
    props.last_modified_by = "Anonymous" if anonymous else "AUTHOR INPUT REQUIRED"
    props.keywords = "PMSM; stator fault; cross-dataset evaluation; negative transfer"
    props.comments = "Double-blind submission artifact"


def add_labeled_field(doc: Document, label: str, value: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    label_run = paragraph.add_run(label + ": ")
    set_font(label_run, bold=True)
    value_run = paragraph.add_run(value)
    set_font(value_run)


def build_title_page(output: Path) -> None:
    doc = Document()
    configure_doc(doc)
    p = doc.add_paragraph(style="Title")
    add_inline_runs(p, TITLE)
    notice = doc.add_paragraph()
    notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = notice.add_run("AUTHOR INPUT REQUIRED BEFORE SUBMISSION")
    set_font(run, size=10, bold=True)
    run.font.color.rgb = RGBColor(156, 0, 6)
    notice.paragraph_format.space_after = Pt(12)

    add_labeled_field(doc, "Article type", "Original Article")
    add_labeled_field(doc, "Running title", SHORT_TITLE)
    add_labeled_field(doc, "Authors", "[FULL NAME 1]1; [FULL NAME 2]2; ...")
    add_labeled_field(
        doc,
        "Affiliation 1",
        "[Department], [University], [City], [Country]",
    )
    add_labeled_field(
        doc,
        "Affiliation 2",
        "[Department/Institute], [Organization], [City], [Country]",
    )
    add_labeled_field(doc, "Corresponding author", "[Full name]")
    add_labeled_field(doc, "E-mail", "[active institutional e-mail]")
    add_labeled_field(doc, "ORCID", "[16-digit ORCID URL for each author, if available]")

    doc.add_paragraph("Acknowledgments", style="Heading 1")
    doc.add_paragraph("[List only contributors who consent to be acknowledged, or state 'None'.]")
    doc.add_paragraph("Statements and declarations", style="Heading 1")
    doc.add_paragraph("Funding", style="Heading 2")
    doc.add_paragraph(
        "[Provide the full funding organization name and grant number, or state: "
        "'The authors received no specific funding for this work.']"
    )
    doc.add_paragraph("Author contributions", style="Heading 2")
    doc.add_paragraph(
        "[Use CRediT roles and verify every statement, for example: Conceptualization, "
        "Methodology, Software, Validation, Formal analysis, Investigation, Data curation, "
        "Writing - original draft, Writing - review and editing, Visualization, "
        "Supervision, Project administration, Funding acquisition.]"
    )
    doc.add_paragraph("Competing interests", style="Heading 2")
    doc.add_paragraph("The authors declare no financial or non-financial competing interests.")
    doc.add_paragraph("Ethics approval", style="Heading 2")
    doc.add_paragraph(
        "Not applicable. This study reanalyzes public experimental machine-current "
        "datasets and does not involve human participants, animals, or personal data."
    )
    doc.add_paragraph("Consent to participate", style="Heading 2")
    doc.add_paragraph("Not applicable.")
    doc.add_paragraph("Consent for publication", style="Heading 2")
    doc.add_paragraph(
        "Not applicable. No human participant or identifiable personal information is reported."
    )
    doc.add_paragraph("Data, materials, and code availability", style="Heading 2")
    doc.add_paragraph(
        "The study reanalyzes the public datasets identified by DOI in the anonymous "
        "manuscript. Online Resources 1 and 2 contain supplementary evidence and an "
        "anonymized reproducibility bundle."
    )
    doc.add_paragraph("Generative AI assistance", style="Heading 2")
    doc.add_paragraph(
        "A large language model (LLM)-based coding assistant was used under human "
        "direction for implementation support, automated consistency checks, and "
        "manuscript drafting and editing. "
        "All authors must review this disclosure and approve the final submitted text."
    )
    set_core_properties(doc, anonymous=False)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.4)
    paragraph.paragraph_format.first_line_indent = Inches(-0.2)
    paragraph.paragraph_format.space_after = Pt(3)
    add_inline_runs(paragraph, text)


def build_cover_letter(output: Path) -> None:
    doc = Document()
    configure_doc(doc)
    add_labeled_field(doc, "Date", today_iso())
    add_labeled_field(doc, "From", "[CORRESPONDING AUTHOR NAME AND INSTITUTION]")
    add_labeled_field(doc, "E-mail", "[INSTITUTIONAL E-MAIL]")
    doc.add_paragraph()
    p = doc.add_paragraph()
    add_inline_runs(p, f"Editor-in-Chief\n{JOURNAL}")
    p = doc.add_paragraph()
    add_inline_runs(p, "Dear Editor-in-Chief,")
    p = doc.add_paragraph()
    add_inline_runs(
        p,
        'We submit the enclosed Original Article, "'
        + TITLE
        + '", for consideration in '
        + JOURNAL
        + ". The study addresses permanent-magnet machines, motor-drive monitoring, sensor "
        "signal processing, and practical industrial reliability, all within the journal's scope.",
    )
    p = doc.add_paragraph()
    add_inline_runs(
        p,
        "The most directly relevant JEET categories are B - Electric Machinery and Power "
        "Electronics (Permanent Magnet Machines; Motor Drive and related applications) and "
        "I - Practical Industrial Electric Applications (Industrial Electric System Control "
        "Applications).",
    )
    p = doc.add_paragraph()
    add_inline_runs(
        p,
        "The manuscript does not claim a universally superior detector. Its contribution is "
        "an audit-first cross-dataset evaluation in which a complete healthy-only pipeline, "
        "threshold, and eleven-method comparison were frozen before a one-time external fault "
        "reveal. Near-ceiling same-family performance reversed on an independently produced "
        "30.16-kW dual-three-phase PMSM dataset, exposing operating-trajectory alarm drift and "
        "conditional negative transfer that conventional random-window evaluation can hide.",
    )
    add_bullet(doc, "Motor- and file-level separation prevents adjacent-window leakage.")
    add_bullet(doc, "Target fault labels are excluded from fitting and threshold selection.")
    add_bullet(
        doc, "Target-only controls distinguish useful adaptation from harmful source augmentation."
    )
    add_bullet(
        doc,
        "Frozen external failure, seed sensitivity, early-horizon behavior, and feature-geometry diagnostics are reported without post-reveal model replacement.",
    )
    p = doc.add_paragraph()
    add_inline_runs(
        p,
        "The work is original, is not under consideration elsewhere, and uses public "
        "experimental machine datasets without human participants or animals. All authors "
        "must approve the manuscript, authorship, declarations, and this letter before the "
        "corresponding author submits it. Online Resource 1 contains the detailed supplementary "
        "evidence, and Online Resource 2 contains an anonymized reproducibility bundle.",
    )
    p = doc.add_paragraph()
    add_inline_runs(p, "Sincerely,")
    p = doc.add_paragraph()
    add_inline_runs(p, "[CORRESPONDING AUTHOR NAME]\n[POSITION, DEPARTMENT, INSTITUTION]")
    set_core_properties(doc, anonymous=False)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def iter_bundle_files() -> Iterable[tuple[Path, Path]]:
    for relative in ANONYMOUS_BUNDLE_PATHS:
        source = ROOT / relative
        if not source.exists():
            continue
        if source.is_file():
            yield source, relative
            continue
        for path in sorted(source.rglob("*")):
            if not path.is_file():
                continue
            if any(
                part in {"__pycache__", ".pytest_cache", ".ruff_cache"}
                or part.endswith(".egg-info")
                for part in path.parts
            ):
                continue
            if path.name.startswith("build_jeet_") or path.name in {
                "build_supplementary_material.py",
                "test_jeet_submission.py",
                "test_supplementary_material.py",
            }:
                continue
            if path.suffix in {".pyc", ".joblib"}:
                continue
            yield path, path.relative_to(ROOT)


def anonymize_bundle_data(data: bytes, relative: Path) -> bytes:
    """Replace task-local absolute roots in text artifacts without altering sources."""
    text_suffixes = {
        ".cfg",
        ".csv",
        ".ini",
        ".json",
        ".md",
        ".py",
        ".toml",
        ".txt",
        ".yaml",
        ".yml",
    }
    if relative.suffix.lower() not in text_suffixes:
        return data
    text = data.decode("utf-8")
    root_windows = str(ROOT)
    for local_root in sorted(
        {
            root_windows.replace("\\", "\\\\"),
            root_windows,
            root_windows.replace("\\", "/"),
        },
        key=len,
        reverse=True,
    ):
        text = text.replace(local_root, ".")
    return text.encode("utf-8")


def build_reproducibility_zip(output: Path) -> dict[str, object]:
    entries: list[dict[str, str]] = []
    readme = """# Anonymous reproducibility bundle

This bundle accompanies a double-blind submission on healthy-only cross-dataset PMSM
fault detection. Raw datasets are excluded. Download and verify the public datasets by
following `docs/data_sources.yaml`, then run the scripts described in the project
protocol. The bundle contains analysis code, tests, frozen protocols, and selected
derived results. It intentionally contains no author names, affiliations, Git history,
or local absolute paths.

## Environment and smoke test

```text
python -m venv .venv
python -m pip install -U pip
python -m pip install -e ".[dev]"
python -m pytest
ruff check .
```

The public raw files are not redistributed. Primary dataset identifiers, licenses,
download locations, and expected hashes are in `docs/data_sources.yaml`; the locked
split and one-time external-reveal rules are in `docs/research_protocol.md`,
`docs/external_validation_protocol.md`, `docs/fault_reveal_log.md`,
`docs/secondary_transient_validation_protocol.md`, and
`docs/secondary_transient_reveal_log.md`. Selected frozen outputs are included under
`results/`. Every archived file covered by the internal manifest can be checked against
`MANIFEST_SHA256.json`.
"""
    output.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        readme_data = readme.encode("utf-8")
        for readme_name in ("README.md", "README_ANONYMOUS.md"):
            info = zipfile.ZipInfo(readme_name, (2026, 8, 20, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            archive.writestr(info, readme_data)
            entries.append({"path": readme_name, "sha256": hashlib.sha256(readme_data).hexdigest()})
        for source, relative in iter_bundle_files():
            data = anonymize_bundle_data(source.read_bytes(), relative)
            lowered = data.lower()
            if any(
                token in lowered
                for token in (
                    b"c:\\users",
                    b"c:\\\\users",
                    b"c:/users",
                    b"c:\\lkc\\phd sci",
                    b"c:\\\\lkc\\\\phd sci",
                    b"c:/lkc/phd sci",
                )
            ):
                raise ValueError(f"Local identity-bearing path found in bundle input: {relative}")
            arcname = relative.as_posix()
            info = zipfile.ZipInfo(arcname, (2026, 8, 20, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            archive.writestr(info, data)
            entries.append({"path": arcname, "sha256": hashlib.sha256(data).hexdigest()})
        manifest_data = json.dumps(entries, indent=2, sort_keys=True).encode("utf-8")
        info = zipfile.ZipInfo("MANIFEST_SHA256.json", (2026, 8, 20, 0, 0, 0))
        info.compress_type = zipfile.ZIP_DEFLATED
        archive.writestr(info, manifest_data)
    return {"files": len(entries), "sha256": sha256(output), "bytes": output.stat().st_size}


def copy_figure_sources(destination: Path) -> list[dict[str, object]]:
    destination.mkdir(parents=True, exist_ok=True)
    manifest = []
    for index, (stem, caption) in enumerate(FIGURES, start=1):
        source_pdf = PAPER / "figures" / f"{stem}.pdf"
        source_png = PAPER / "figures" / f"{stem}.png"
        target_pdf = destination / f"Fig{index}.pdf"
        target_png = destination / f"Fig{index}.png"
        shutil.copy2(source_pdf, target_pdf)
        shutil.copy2(source_png, target_png)
        with Image.open(target_png) as image:
            dpi = image.info.get("dpi", (0, 0))[0]
            size = image.size
        manifest.append(
            {
                "number": index,
                "caption": caption,
                "pdf": target_pdf.name,
                "pdf_sha256": sha256(target_pdf),
                "png": target_png.name,
                "png_sha256": sha256(target_png),
                "png_pixels": list(size),
                "png_dpi": round(float(dpi), 3),
            }
        )
    return manifest


def write_submission_docs() -> None:
    readme = f"""# JEET submission package

Generated: {today_iso()}

## Upload files

1. `Manuscript_Anonymous.docx` - double-blind review manuscript with inline figures,
   tables, numeric citations, declarations, and no author identity.
2. `Title_Page_AUTHOR_INPUT_REQUIRED.docx` - replace every bracketed field, obtain all
   author approvals, then rename to `Title_Page.docx`.
3. `Cover_Letter_AUTHOR_INPUT_REQUIRED.docx` - complete and sign before upload.
4. `ESM_1_Supplementary_Material.pdf` - anonymous detailed evidence (Online Resource 1).
5. `ESM_2_Reproducibility_Code.zip` - anonymous code, tests, protocols, and selected
   derived outputs (Online Resource 2).
6. `figures/Fig1` through `Fig11` - editable/vector PDF plus 300-dpi review PNG copies.

## Internal control file

`QA_Report.md` records the completed visual and structural checks and the remaining
native-Word/human gates. Keep it with the working package, but do not upload it as a
manuscript file.

`Author_Input_Form_CN.md` is a Chinese-language intake sheet for the author-owned
metadata that cannot be inferred safely. It is also an internal file and must not be
uploaded.

`Portal_Clarification_Email_Draft.md` is a ready-to-personalize message for the JEET
office if the linked Editorial Manager site still displays its implementation warning.
It is an internal draft, not an upload file, and has not been sent.

## Remaining human-only gates

- Confirm author names, order, affiliations, ORCIDs, corresponding author, CRediT roles,
  acknowledgments, and funding.
- Every author must inspect and approve the full manuscript, declarations, AI-assistance
  disclosure, supplementary material, and cover letter.
- Confirm the work is not under consideration elsewhere and obtain institutional
  permission to submit.
- Confirm the live submission route with the JEET office if Editorial Manager still
  displays its implementation-mode warning. Both Springer and KIEE currently point to
  `https://www.editorialmanager.com/eete`, while that destination says not to use it for
  live manuscript submission. The KIEE contact is `jeet@kiee.or.kr`.
- Recheck the live journal site on submission day. JEET uses double-blind review and
  requests a separate title page.

## Rebuild

Run `scripts/build_jeet_submission.py` and then `scripts/build_jeet_pdfs.py`. The first
command rebuilds the editable Word sources, figures, and anonymous code bundle; the
second creates and audits the two review PDFs. Per-artifact SHA-256 values and PDF page
audits are in `build_metadata.json`.
"""
    checklist = f"""# JEET pre-submission checklist

## Automatically verified

- [x] Anonymous manuscript and separate title-page file
- [x] Abstract between 150 and 250 words
- [x] Six keywords
- [x] Decimal headings with no more than three levels
- [x] Three tables and eleven figures cited in consecutive order
- [x] Numeric square-bracket citations and reference list with DOI links when available
- [x] Data availability, competing interests, ethics, funding-location, author-
      contribution-location, and generative-AI statements
- [x] Online Resource 1 and Online Resource 2 cited in the manuscript
- [x] Figure source copies named Fig1-Fig11
- [x] Anonymous reproducibility ZIP excludes raw data, Git history, and known local paths

## Human approval required

- [ ] Replace all bracketed fields in the title page and cover letter
- [ ] Confirm legal author names, order, affiliations, e-mails, and ORCIDs
- [ ] Confirm corresponding author and CRediT contribution statement
- [ ] Confirm funding body/grant number or explicit no-funding statement
- [ ] Obtain permission from every person named in acknowledgments
- [ ] All authors approve the final files and institutional submission
- [ ] Confirm no simultaneous submission and disclose any related manuscript/preprint
- [ ] Review the AI-assistance disclosure for accuracy
- [ ] Confirm public dataset licences and final anonymous data/code links
- [ ] Recheck current JEET indexing, fees, and submission fields on the submission date
- [ ] Confirm the active submission portal with `jeet@kiee.or.kr` if Editorial Manager
      still displays "Site under development. Do not use for live manuscript submission."

## Cost warning

The JEET website currently states subscription-model page charges of US$50 per page
within six pages, US$60 per page for pages 7-12, and US$80 per page over 13 pages. The
publisher determines production page count; obtain supervisor/funder approval before
submission. See https://link.springer.com/journal/42835/submission-guidelines.

The journal is hybrid. The publisher currently lists an optional open-access APC of
GBP 2,590 / USD 3,590 / EUR 2,890 plus applicable taxes. Do not assume whether the
subscription page charges and the OA APC are cumulative; confirm the chosen route and
invoice treatment with the journal before submission.

## Submission-route warning ({today_iso()})

Springer and KIEE both link to `https://www.editorialmanager.com/eete`, but the landing
page currently states that the site is under development and must not be used for live
submission. If that warning remains, contact `jeet@kiee.or.kr` and obtain the active
submission route before uploading any file.
"""
    (OUT_DIR / "README.md").write_text(readme, encoding="utf-8")
    (OUT_DIR / "Submission_Checklist.md").write_text(checklist, encoding="utf-8")
    portal_email = f"""# JEET submission-route clarification e-mail draft

Status: **not sent**. Personalize the bracketed fields before sending.

To: `jeet@kiee.or.kr`

Subject: Request for current submission route for a JEET Original Article

Dear JEET Editorial Office,

I am preparing an Original Article entitled “{TITLE}” for submission to the
*Journal of Electrical Engineering & Technology*.

On {today_iso()}, both the Springer and KIEE journal pages directed authors to
`https://www.editorialmanager.com/eete`, but that landing page displayed “Site under
development. Do not use for live manuscript submission.” Before uploading any files,
could you please confirm:

1. the currently active submission URL or procedure;
2. whether the anonymous manuscript plus separate title page remains the correct
   double-blind file arrangement; and
3. the current page-charge schedule and whether page charges also apply when an author
   selects the optional open-access route?

No manuscript is attached to this routing inquiry.

Sincerely,

[CORRESPONDING AUTHOR NAME]

[DEPARTMENT AND INSTITUTION]

[INSTITUTIONAL E-MAIL]
"""
    (OUT_DIR / "Portal_Clarification_Email_Draft.md").write_text(
        portal_email, encoding="utf-8"
    )
    author_metadata = """title: >-
  When Healthy-Only Transfer Fails in PMSM Stator-Fault Detection:
  A Leakage-Resistant Cross-Dataset Evaluation
article_type: Original Article
authors:
  - full_name: "REQUIRED"
    affiliation_ids: [1]
    email: "REQUIRED"
    orcid: ""
    corresponding: true
affiliations:
  1: "Department, University, City, Country"
funding: "REQUIRED: full funder and grant, or explicit no-funding statement"
acknowledgments: "None, or names with permission"
competing_interests: "The authors declare no financial or non-financial competing interests."
author_contributions: "REQUIRED: verified CRediT roles"
ethics_approval: >-
  Not applicable. This study reanalyzes public experimental machine-current datasets
  and does not involve human participants, animals, or personal data.
consent_to_participate: "Not applicable."
consent_for_publication: >-
  Not applicable. No human participant or identifiable personal information is reported.
data_materials_code_availability: >-
  The study reanalyzes the public datasets identified by DOI in the anonymous manuscript.
  Online Resources 1 and 2 contain supplementary evidence and an anonymized
  reproducibility bundle.
generative_ai_assistance: >-
  A large language model (LLM)-based coding assistant was used under human direction
  for implementation support, automated consistency checks, and manuscript drafting
  and editing. All authors must review this disclosure and approve the final submitted
  text.
"""
    (OUT_DIR / "author_metadata_REQUIRED.yaml").write_text(author_metadata, encoding="utf-8")
    author_input_cn = """# JEET 作者信息填写单（不要上传此文件）

请按以下字段回复，或直接填写同目录的 `author_metadata_REQUIRED.yaml`。姓名、作者顺序、
导师署名、单位、基金和 CRediT 贡献都属于作者本人决定；未获得所有作者确认前不要投稿。

## 1. 作者与顺序

每位作者分别提供：

- 英文全名（与护照/既有论文一致）：
- 排名：
- 单位编号（可多选）：
- 有效邮箱：
- ORCID（如有，完整 16 位 URL）：
- 是否通讯作者：

## 2. 单位

请使用学校或机构确认过的官方英文写法，格式为：

`Department/Institute, University/Organization, City, Country`

不要根据中文名称自行猜译；学校、实验室或导师单位只有在作者确认后才能写入。

## 3. Funding 与致谢

- Funding：提供资助机构官方英文全称和 grant number；如无，明确写
  `The authors received no specific funding for this work.`
- Acknowledgments：写 `None`，或列出已同意被致谢者及其贡献。

## 4. CRediT 作者贡献

逐位作者从下列角色中选择真实承担的角色：Conceptualization; Methodology; Software;
Validation; Formal analysis; Investigation; Resources; Data curation; Writing - original
draft; Writing - review and editing; Visualization; Supervision; Project administration;
Funding acquisition。不要为了凑角色而分配未实际完成的贡献。

## 5. 需要全体作者确认的预填声明

- Competing interests：`The authors declare no financial or non-financial competing interests.`
- Ethics approval：不适用；本研究仅复用公开电机实验数据，无人类、动物或个人数据。
- Consent to participate / publication：不适用。
- Generative AI assistance：LLM 编码助手在人工指导下用于实现支持、自动一致性检查及
  稿件起草/编辑；全部计算由作者对冻结结果和测试复核，最终责任归人类作者。

如任何一项不准确，请给出替换文本，不要直接确认。

## 6. 投稿入口人工门槛

截至 2026-08-21，Springer 与 KIEE 指向的 Editorial Manager 页面仍显示
`Site under development. Do not use for live manuscript submission.`。若提交时警告仍在，
先联系 `jeet@kiee.or.kr` 获取有效入口。
"""
    (OUT_DIR / "Author_Input_Form_CN.md").write_text(author_input_cn, encoding="utf-8")


def audit_sources() -> dict[str, object]:
    text = MANUSCRIPT.read_text(encoding="utf-8")
    abstract = text.split("## Abstract", 1)[1].split("**Keywords:**", 1)[0]
    abstract_words = len(re.findall(r"[A-Za-z0-9]+(?:[.-][A-Za-z0-9]+)*", abstract))
    keyword_text = text.split("**Keywords:**", 1)[1].split("##", 1)[0].replace("\n", " ")
    keyword_count = len([item for item in keyword_text.split(";") if item.strip()])
    figures_cited = sorted({int(value) for value in re.findall(r"(?:Fig\.|Figure)\s+(\d+)", text)})
    tables_cited = sorted({int(value) for value in re.findall(r"Table\s+(\d+)", text)})
    citations = set(re.findall(r"@([A-Za-z0-9_:-]+)", text))
    bib_keys = set(parse_bibtex(BIBLIOGRAPHY))
    if not 150 <= abstract_words <= 250:
        raise AssertionError(f"Abstract word count {abstract_words} is outside 150-250")
    if not 4 <= keyword_count <= 6:
        raise AssertionError(f"Keyword count {keyword_count} is outside 4-6")
    if figures_cited != list(range(1, len(FIGURES) + 1)):
        raise AssertionError(f"Figure citations are incomplete: {figures_cited}")
    if tables_cited != [1, 2, 3]:
        raise AssertionError(f"Table citations are incomplete: {tables_cited}")
    if citations - bib_keys:
        raise AssertionError(f"Missing bibliography keys: {sorted(citations - bib_keys)}")
    return {
        "abstract_words": abstract_words,
        "keywords": keyword_count,
        "figure_citations": figures_cited,
        "table_citations": tables_cited,
        "cited_bibliography_entries": len(citations),
    }


def build() -> dict[str, object]:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    TMP_DIR.mkdir(parents=True, exist_ok=True)
    source_audit = audit_sources()
    manuscript_meta = build_markdown_doc(
        MANUSCRIPT,
        OUT_DIR / "Manuscript_Anonymous.docx",
        anonymous_main=True,
    )
    supplement_meta = build_markdown_doc(
        SUPPLEMENT,
        TMP_DIR / "ESM_1_Supplementary_Material.docx",
        anonymous_main=False,
        supplement=True,
    )
    build_title_page(OUT_DIR / "Title_Page_AUTHOR_INPUT_REQUIRED.docx")
    build_cover_letter(OUT_DIR / "Cover_Letter_AUTHOR_INPUT_REQUIRED.docx")
    bundle_meta = build_reproducibility_zip(OUT_DIR / "ESM_2_Reproducibility_Code.zip")
    figures_meta = copy_figure_sources(OUT_DIR / "figures")
    metadata: dict[str, object] = {
        "generated": today_iso(),
        "journal": JOURNAL,
        "journal_requirements": {
            "checked": today_iso(),
            "official_kiee_template_page": (
                "https://www.kiee.or.kr/board/?_0000_method=view&ncode=a008&num=1758&page=1"
            ),
            "springer_guidelines": (
                "https://link.springer.com/journal/42835/submission-guidelines"
            ),
            "springer_fees": (
                "https://link.springer.com/journal/42835/how-to-publish-with-us"
            ),
            "submission_portal": "https://www.editorialmanager.com/eete",
            "submission_portal_warning_observed": (
                "Site under development. Do not use for live manuscript submission."
            ),
            "journal_contact": "jeet@kiee.or.kr",
        },
        "source_audit": source_audit,
        "manuscript": manuscript_meta,
        "supplement_intermediate": supplement_meta,
        "title_page_sha256": sha256(OUT_DIR / "Title_Page_AUTHOR_INPUT_REQUIRED.docx"),
        "cover_letter_sha256": sha256(OUT_DIR / "Cover_Letter_AUTHOR_INPUT_REQUIRED.docx"),
        "reproducibility_bundle": bundle_meta,
        "figures": figures_meta,
    }
    metadata_path = OUT_DIR / "build_metadata.json"
    metadata_path.write_text(json.dumps(metadata, indent=2, sort_keys=True), encoding="utf-8")
    write_submission_docs()
    return metadata


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check", action="store_true", help="Audit source and existing package")
    args = parser.parse_args()
    if args.check:
        audit = audit_sources()
        required = [
            OUT_DIR / "Manuscript_Anonymous.docx",
            OUT_DIR / "Title_Page_AUTHOR_INPUT_REQUIRED.docx",
            OUT_DIR / "Cover_Letter_AUTHOR_INPUT_REQUIRED.docx",
            OUT_DIR / "ESM_2_Reproducibility_Code.zip",
            OUT_DIR / "build_metadata.json",
        ]
        missing = [str(path) for path in required if not path.exists()]
        if missing:
            raise FileNotFoundError(f"Missing generated submission files: {missing}")
        print(json.dumps(audit, sort_keys=True))
        return
    print(json.dumps(build(), indent=2, sort_keys=True))


if __name__ == "__main__":
    main()

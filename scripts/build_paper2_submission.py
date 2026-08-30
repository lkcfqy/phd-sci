"""Build journal-neutral Paper 2 DOCX and PDF manuscripts from the frozen Markdown."""

from __future__ import annotations

import argparse
import hashlib
import html
import json
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image as PILImage
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    Image,
    KeepTogether,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

import build_jeet_submission as legacy

ROOT = Path(__file__).resolve().parents[1]
PAPER_DIR = ROOT / "papers" / "paper2_torque_uq"
MANUSCRIPT = PAPER_DIR / "manuscript.md"
BIBLIOGRAPHY = ROOT / "references" / "key_papers.bib"
OUT_DIR = PAPER_DIR / "submission"
DOCX_OUT = OUT_DIR / "Paper2_Manuscript_Anonymous.docx"
PDF_OUT = OUT_DIR / "Paper2_Manuscript_Anonymous.pdf"
METADATA_OUT = OUT_DIR / "build_metadata.json"
SHORT_TITLE = "Curvewise Conformal PMSM Torque Bands"
ARTIFACT_LABEL = "Paper 2"
DOCUMENT_SUBJECT = "Journal-neutral Paper 2 submission draft"
EXPECTED_FIGURES = 5
EXPECTED_TABLES = 5


@dataclass
class Block:
    kind: str
    text: str = ""
    level: int = 0
    ordinal: int = 0
    rows: list[list[str]] = field(default_factory=list)
    image_path: Path | None = None
    alt_text: str = ""


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalize_text(value: str) -> str:
    return (
        value.replace("–", "-")
        .replace("—", "-")
        .replace("−", "-")
        .replace("‑", "-")
        .replace("--", "-")
        .replace(" ", " ")
    )


def parse_markdown(path: Path) -> tuple[str, list[Block]]:
    lines = path.read_text(encoding="utf-8").splitlines()
    title = "Paper 2 manuscript"
    index = 0
    if lines and lines[0].strip() == "---":
        end = next(i for i in range(1, len(lines)) if lines[i].strip() == "---")
        for line in lines[1:end]:
            match = re.match(r'title:\s*["\']?(.*?)["\']?\s*$', line.strip())
            if match:
                title = match.group(1)
                break
        index = end + 1

    blocks: list[Block] = []

    def starts_block(line: str) -> bool:
        stripped = line.strip()
        return bool(
            not stripped
            or stripped.startswith(("#", "!", "|", "- "))
            or re.match(r"^\d+\.\s+", stripped)
            or stripped == "$$"
        )

    while index < len(lines):
        stripped = lines[index].strip()
        if not stripped:
            index += 1
            continue
        if stripped == "$$":
            index += 1
            equation_lines: list[str] = []
            while index < len(lines) and lines[index].strip() != "$$":
                equation_lines.append(lines[index].strip())
                index += 1
            if index >= len(lines):
                raise ValueError("Unclosed display equation")
            index += 1
            blocks.append(Block("equation", text=" ".join(equation_lines)))
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", stripped)
        if heading:
            blocks.append(
                Block("heading", text=heading.group(2), level=len(heading.group(1)))
            )
            index += 1
            continue
        image = re.match(r"^!\[(.*?)\]\((.*?)\)$", stripped)
        if image:
            image_path = (path.parent / image.group(2)).resolve()
            blocks.append(
                Block(
                    "image",
                    image_path=image_path,
                    alt_text=image.group(1),
                )
            )
            index += 1
            continue
        if stripped.startswith("|"):
            rows: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = [
                    cell.replace(r"\|", "|").strip()
                    for cell in re.split(r"(?<!\\)\|", lines[index].strip())[1:-1]
                ]
                separator = all(
                    re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells
                )
                if not separator:
                    rows.append(cells)
                index += 1
            blocks.append(Block("table", rows=rows))
            continue
        list_match = re.match(r"^(?P<number>\d+)\.\s+(?P<text>.+)$", stripped)
        if stripped.startswith("- ") or list_match:
            list_kind = "number" if list_match else "bullet"
            item = list_match.group("text") if list_match else stripped[2:]
            index += 1
            continuations = [item]
            while index < len(lines) and not starts_block(lines[index]):
                continuations.append(lines[index].strip())
                index += 1
            blocks.append(
                Block(
                    list_kind,
                    text=" ".join(continuations),
                    ordinal=int(list_match.group("number")) if list_match else 0,
                )
            )
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines) and not starts_block(lines[index]):
            paragraph_lines.append(lines[index].strip())
            index += 1
        blocks.append(Block("paragraph", text=" ".join(paragraph_lines)))
    return normalize_text(title), blocks


DISPLAY_EQUATIONS = {
    "s_i=\\max_{j\\in\\{0,\\ldots,119\\}}\\left|y_{ij}-\\haty_{ij}\\right|.": (
        "sᵢ = max(j in {0,…,119}) |yᵢⱼ − ŷᵢⱼ|"
    ),
    "q_\\alpha=s_{(k)},\\qquadk=\\left\\lceil(n_{cal}+1)(1-\\alpha)\\right\\rceil.": (
        "q(α) = s₍ₖ₎,    k = ceil[(ncal + 1)(1 − α)]"
    ),
    (
        "C_{global}(x,j)=[\\haty_j(x)-q_\\alpha,\\;\\haty_j(x)+q_\\alpha],"
        "\\qquadj=0,\\ldots,119."
    ): "Cglobal(x,j) = [ŷⱼ(x) − q(α), ŷⱼ(x) + q(α)],    j = 0,…,119",
    "g(x)=\\max\\left\\{0.25,\\frac{d_5(x;X_{fit})}{m_5}\\right\\}.": (
        "g(x) = max{0.25, d₅(x; Xfit) / m₅}"
    ),
    (
        "C_{scaled}(x,j)=[\\haty_j(x)-q^g_\\alphag(x),\\;"
        "\\haty_j(x)+q^g_\\alphag(x)]."
    ): (
        "Cscaled(x,j) = [ŷⱼ(x) − qᵍ(α)g(x), "
        "ŷⱼ(x) + qᵍ(α)g(x)]"
    ),
    (
        "p_{support}(x_*)=\\frac{1+\\sum_{i=1}^{600}\\mathbb{1}"
        "\\{d_i\\ged_*\\}}{601}."
    ): "psupport(x*) = [1 + Σᵢ₌₁⁶⁰⁰ 1{dᵢ ≥ d*}] / 601",
    "ESS=\\frac{(\\sum_iw_i)^2}{\\sum_iw_i^2},": (
        "ESS = (Σᵢ wᵢ)² / Σᵢ wᵢ²"
    ),
}


def display_equation(value: str) -> str:
    canonical = re.sub(r"\s+", "", value)
    if canonical in DISPLAY_EQUATIONS:
        return DISPLAY_EQUATIONS[canonical]
    return inline_math(value)


def inline_math(value: str) -> str:
    text = value
    replacements = {
        r"\mathbb{R}": "ℝ",
        r"\mathbb{1}": "1",
        r"\hat y": "ŷ",
        r"\haty": "ŷ",
        r"\alpha": "α",
        r"\ge": "≥",
        r"\le": "≤",
        r"\in": " in ",
        r"\ldots": "…",
        r"\max": "max",
        r"\sum": "sum",
        r"\frac": "frac",
        r"\left": "",
        r"\right": "",
        r"\qquad": "    ",
        r"\;": " ",
        r"\,": " ",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    text = re.sub(r"frac\{([^{}]+)\}\{([^{}]+)\}", r"(\1)/(\2)", text)
    text = re.sub(r"\\mathrm\{([^{}]+)\}", r"\1", text)
    text = re.sub(r"\\operatorname\{([^{}]+)\}", r"\1", text)
    text = text.replace("{", "").replace("}", "").replace("\\", "")
    return re.sub(r"\s+", " ", text).strip()


def replace_citations(value: str, cite_state: dict[str, int]) -> str:
    return legacy.replace_citations(value, cite_state)


def clean_reference(value: str) -> str:
    """Finish the small TeX-accent subset not decoded by the legacy formatter."""
    replacements = {
        r"\`a": "à",
        r"\`e": "è",
        r"\`i": "ì",
        r"\`o": "ò",
        r"\`u": "ù",
        r"\'a": "á",
        r"\'e": "é",
        r"\'i": "í",
        r"\'o": "ó",
        r"\'u": "ú",
        r'\"a': "ä",
        r'\"o': "ö",
        r'\"u': "ü",
    }
    for source, target in replacements.items():
        value = value.replace(source, target)
    return value


def set_run_font(
    run,
    *,
    name: str = "Times New Roman",
    size: float = 10.5,
    bold: bool | None = None,
    italic: bool | None = None,
    color: RGBColor | None = None,
) -> None:
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_docx_inline(paragraph, value: str, cite_state: dict[str, int]) -> None:
    value = normalize_text(replace_citations(value, cite_state))
    token_re = re.compile(
        r"(\\\(.+?\\\)|\*\*.+?\*\*|`.+?`|(?<!\*)\*[^*]+?\*)"
    )
    cursor = 0
    for match in token_re.finditer(value):
        if match.start() > cursor:
            set_run_font(paragraph.add_run(value[cursor : match.start()]))
        token = match.group(0)
        if token.startswith(r"\("):
            set_run_font(
                paragraph.add_run(inline_math(token[2:-2])),
                name="Cambria Math",
                italic=True,
            )
        elif token.startswith("**"):
            set_run_font(paragraph.add_run(token[2:-2]), bold=True)
        elif token.startswith("`"):
            set_run_font(
                paragraph.add_run(token[1:-1]), name="Courier New", size=9.0
            )
        else:
            set_run_font(paragraph.add_run(token[1:-1]), italic=True)
        cursor = match.end()
    if cursor < len(value):
        set_run_font(paragraph.add_run(value[cursor:]))


def add_page_number(section) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    cached = OxmlElement("w:t")
    cached.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, cached, end))
    set_run_font(run, size=8.5, color=RGBColor(90, 90, 90))


def configure_docx(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(
        header.add_run(SHORT_TITLE),
        size=8.5,
        italic=True,
        color=RGBColor(100, 100, 100),
    )
    add_page_number(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.2

    title = doc.styles["Title"]
    title.font.name = "Times New Roman"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    title.font.size = Pt(16)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(12)
    title.paragraph_format.space_after = Pt(8)

    heading_specs = {
        "Heading 1": (13.0, 14, 6, False),
        "Heading 2": (11.5, 12, 5, False),
        "Heading 3": (10.5, 10, 4, True),
    }
    for name, (size, before, after, italic) in heading_specs.items():
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.italic = italic
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_style in ("List Bullet", "List Number"):
        style = doc.styles[list_style]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.194)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.208

    if "Paper Caption" not in doc.styles:
        doc.styles.add_style("Paper Caption", 1)
    caption = doc.styles["Paper Caption"]
    caption.font.name = "Times New Roman"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    caption.font.size = Pt(9)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(7)
    caption.paragraph_format.keep_with_next = False


def add_docx_table(doc: Document, rows: list[list[str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    widths = legacy.column_widths(rows, 9360)
    legacy.set_table_geometry(table, widths, indent_dxa=120)
    legacy.set_repeat_table_header(table.rows[0])
    for row_index, (source_row, target_row) in enumerate(zip(rows, table.rows, strict=True)):
        for column_index, (value, cell) in enumerate(
            zip(source_row, target_row.cells, strict=True)
        ):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                legacy.shade_cell(cell, "F2F4F7")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT
                if column_index == 0
                else WD_ALIGN_PARAGRAPH.CENTER
            )
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            clean = normalize_text(value.replace("**", "").replace("`", ""))
            set_run_font(
                paragraph.add_run(clean),
                size=7.6,
                bold=(row_index == 0),
            )
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)


def add_docx_image(doc: Document, block: Block) -> None:
    if block.image_path is None or not block.image_path.is_file():
        raise FileNotFoundError(block.image_path)
    with PILImage.open(block.image_path) as image:
        width_px, height_px = image.size
    width_in = 6.35
    height_in = width_in * height_px / width_px
    if height_in > 7.25:
        height_in = 7.25
        width_in = height_in * width_px / height_px
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_with_next = True
    shape = paragraph.add_run().add_picture(str(block.image_path), width=Inches(width_in))
    shape._inline.docPr.set("descr", block.alt_text)


def add_docx_references(doc: Document, cite_state: dict[str, int]) -> None:
    entries = legacy.parse_bibtex(BIBLIOGRAPHY)
    missing = sorted(set(cite_state) - set(entries))
    if missing:
        raise KeyError(f"Missing bibliography entries: {missing}")
    doc.add_paragraph("References", style="Heading 1")
    for key, number in sorted(cite_state.items(), key=lambda item: item[1]):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        paragraph.paragraph_format.space_after = Pt(4)
        set_run_font(
            paragraph.add_run(
                f"[{number}] {clean_reference(legacy.format_reference(entries[key]))}"
            ),
            size=9.0,
        )


def build_docx(title: str, blocks: list[Block]) -> dict[str, Any]:
    doc = Document()
    configure_docx(doc)
    title_paragraph = doc.add_paragraph(style="Title")
    add_docx_inline(title_paragraph, title, {})
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)
    set_run_font(
        subtitle.add_run("Anonymous review manuscript | author details withheld"),
        size=9.5,
        italic=True,
        color=RGBColor(80, 80, 80),
    )

    cite_state: dict[str, int] = {}
    figure_count = 0
    table_count = 0
    for block in blocks:
        if block.kind == "heading":
            style = (
                "Heading 1"
                if block.level == 1
                else ("Heading 2" if block.level == 2 else "Heading 3")
            )
            paragraph = doc.add_paragraph(style=style)
            add_docx_inline(paragraph, block.text, cite_state)
        elif block.kind == "paragraph":
            caption = block.text.startswith(("**Figure ", "**Table "))
            paragraph = doc.add_paragraph(style="Paper Caption" if caption else None)
            add_docx_inline(paragraph, block.text, cite_state)
        elif block.kind in {"bullet", "number"}:
            paragraph = doc.add_paragraph(
                style="List Bullet" if block.kind == "bullet" else None
            )
            if block.kind == "number":
                paragraph.paragraph_format.left_indent = Inches(0.375)
                paragraph.paragraph_format.first_line_indent = Inches(-0.194)
                paragraph.paragraph_format.space_after = Pt(4)
                paragraph.paragraph_format.line_spacing = 1.208
                set_run_font(paragraph.add_run(f"{block.ordinal}. "))
            add_docx_inline(paragraph, block.text, cite_state)
        elif block.kind == "equation":
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(6)
            set_run_font(
                paragraph.add_run(display_equation(block.text)),
                name="Cambria Math",
                size=10.5,
            )
        elif block.kind == "image":
            add_docx_image(doc, block)
            figure_count += 1
        elif block.kind == "table":
            add_docx_table(doc, block.rows)
            table_count += 1
        else:
            raise ValueError(f"Unsupported block: {block.kind}")
    add_docx_references(doc, cite_state)

    properties = doc.core_properties
    properties.title = title
    properties.subject = DOCUMENT_SUBJECT
    properties.author = "Anonymous"
    properties.last_modified_by = "Anonymous"
    properties.keywords = "PMSM; conformal prediction; surrogate; torque"
    properties.comments = "Generated from evidence-validated Markdown"
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_OUT)
    return {
        "path": str(DOCX_OUT.relative_to(ROOT)),
        "sha256": sha256(DOCX_OUT),
        "figures": figure_count,
        "tables": table_count,
        "citations": len(cite_state),
        "paragraphs": len(doc.paragraphs),
    }


def register_pdf_fonts() -> None:
    font_dir = Path("C:/Windows/Fonts")
    fonts = {
        "TNR": "times.ttf",
        "TNR-Bold": "timesbd.ttf",
        "TNR-Italic": "timesi.ttf",
        "TNR-BoldItalic": "timesbi.ttf",
    }
    for name, filename in fonts.items():
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(font_dir / filename)))
    pdfmetrics.registerFontFamily(
        "TNR",
        normal="TNR",
        bold="TNR-Bold",
        italic="TNR-Italic",
        boldItalic="TNR-BoldItalic",
    )


def pdf_inline(value: str, cite_state: dict[str, int]) -> str:
    value = normalize_text(replace_citations(value, cite_state))
    placeholders: dict[str, str] = {}

    def protect_math(match: re.Match[str]) -> str:
        key = f"MATHPLACEHOLDER{len(placeholders)}"
        placeholders[key] = html.escape(inline_math(match.group(1)))
        return key

    value = re.sub(r"\\\((.+?)\\\)", protect_math, value)
    value = html.escape(value)
    value = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", value)
    value = re.sub(r"`(.+?)`", r"<font name='Courier'>\1</font>", value)
    value = re.sub(r"(?<!\*)\*([^*]+?)\*", r"<i>\1</i>", value)
    for key, math_text in placeholders.items():
        value = value.replace(key, f"<i>{math_text}</i>")
    return value


def pdf_styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "PaperTitle",
            parent=base["Title"],
            fontName="TNR-Bold",
            fontSize=16,
            leading=19,
            alignment=TA_CENTER,
            textColor=colors.black,
            spaceBefore=10,
            spaceAfter=8,
        ),
        "subtitle": ParagraphStyle(
            "PaperSubtitle",
            parent=base["Normal"],
            fontName="TNR-Italic",
            fontSize=9.5,
            leading=12,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#505050"),
            spaceAfter=12,
        ),
        "body": ParagraphStyle(
            "PaperBody",
            parent=base["BodyText"],
            fontName="TNR",
            fontSize=10.5,
            leading=13.2,
            alignment=TA_JUSTIFY,
            textColor=colors.black,
            spaceAfter=6,
        ),
        "h1": ParagraphStyle(
            "PaperH1",
            parent=base["Heading1"],
            fontName="TNR-Bold",
            fontSize=13,
            leading=16,
            textColor=colors.black,
            spaceBefore=14,
            spaceAfter=6,
            keepWithNext=True,
        ),
        "h2": ParagraphStyle(
            "PaperH2",
            parent=base["Heading2"],
            fontName="TNR-Bold",
            fontSize=11.5,
            leading=14,
            textColor=colors.black,
            spaceBefore=12,
            spaceAfter=5,
            keepWithNext=True,
        ),
        "h3": ParagraphStyle(
            "PaperH3",
            parent=base["Heading3"],
            fontName="TNR-BoldItalic",
            fontSize=10.5,
            leading=13,
            textColor=colors.black,
            spaceBefore=10,
            spaceAfter=4,
            keepWithNext=True,
        ),
        "caption": ParagraphStyle(
            "PaperCaption",
            parent=base["BodyText"],
            fontName="TNR",
            fontSize=9,
            leading=11,
            alignment=TA_LEFT,
            textColor=colors.black,
            spaceBefore=3,
            spaceAfter=7,
        ),
        "equation": ParagraphStyle(
            "PaperEquation",
            parent=base["BodyText"],
            fontName="TNR",
            fontSize=10.5,
            leading=14,
            alignment=TA_CENTER,
            spaceBefore=4,
            spaceAfter=6,
        ),
        "bullet": ParagraphStyle(
            "PaperBullet",
            parent=base["BodyText"],
            fontName="TNR",
            fontSize=10.5,
            leading=12.7,
            leftIndent=0.375 * inch,
            firstLineIndent=-0.194 * inch,
            bulletIndent=0.181 * inch,
            spaceAfter=4,
        ),
        "reference": ParagraphStyle(
            "PaperReference",
            parent=base["BodyText"],
            fontName="TNR",
            fontSize=9,
            leading=11,
            leftIndent=0.25 * inch,
            firstLineIndent=-0.25 * inch,
            spaceAfter=4,
        ),
    }


def pdf_column_widths(rows: list[list[str]], total: float) -> list[float]:
    weights: list[float] = []
    for column in range(len(rows[0])):
        maximum = max(len(re.sub(r"[*`]", "", row[column])) for row in rows)
        weights.append(max(4.0, min(26.0, maximum**0.5 * 2.3)))
    total_weight = sum(weights)
    return [total * weight / total_weight for weight in weights]


def make_pdf_table(
    rows: list[list[str]], styles: dict[str, ParagraphStyle], width: float
) -> Table:
    cells: list[list[Paragraph]] = []
    empty_citations: dict[str, int] = {}
    for row_index, row in enumerate(rows):
        row_cells: list[Paragraph] = []
        for value in row:
            style = ParagraphStyle(
                f"Cell{row_index}",
                parent=styles["body"],
                fontName="TNR-Bold" if row_index == 0 else "TNR",
                fontSize=7.4,
                leading=8.8,
                alignment=TA_CENTER,
                spaceAfter=0,
            )
            row_cells.append(
                Paragraph(
                    pdf_inline(value.replace("**", ""), empty_citations),
                    style,
                )
            )
        cells.append(row_cells)
    table = Table(
        cells,
        colWidths=pdf_column_widths(rows, width),
        repeatRows=1,
        hAlign="LEFT",
        splitByRow=1,
    )
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#777777")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F2F4F7")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return table


def make_pdf_image(block: Block, width: float) -> Image:
    if block.image_path is None or not block.image_path.is_file():
        raise FileNotFoundError(block.image_path)
    with PILImage.open(block.image_path) as image:
        width_px, height_px = image.size
    draw_width = min(width, 6.35 * inch)
    draw_height = draw_width * height_px / width_px
    if draw_height > 7.25 * inch:
        draw_height = 7.25 * inch
        draw_width = draw_height * width_px / height_px
    image = Image(str(block.image_path), width=draw_width, height=draw_height)
    image.hAlign = "CENTER"
    return image


class PaperDocTemplate(BaseDocTemplate):
    def __init__(self, filename: str):
        super().__init__(
            filename,
            pagesize=letter,
            leftMargin=inch,
            rightMargin=inch,
            topMargin=inch,
            bottomMargin=inch,
            title=SHORT_TITLE,
            author="Anonymous",
            subject=DOCUMENT_SUBJECT,
        )
        frame = Frame(
            self.leftMargin,
            self.bottomMargin,
            self.width,
            self.height,
            id="body",
        )
        self.addPageTemplates(PageTemplate(id="paper", frames=[frame], onPage=self._page))

    @staticmethod
    def _page(canvas, document) -> None:
        canvas.saveState()
        canvas.setFont("TNR-Italic", 8.5)
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.drawCentredString(letter[0] / 2, letter[1] - 0.55 * inch, SHORT_TITLE)
        canvas.setFont("TNR", 8.5)
        canvas.drawCentredString(letter[0] / 2, 0.52 * inch, str(document.page))
        canvas.restoreState()


def build_pdf(title: str, blocks: list[Block]) -> dict[str, Any]:
    register_pdf_fonts()
    styles = pdf_styles()
    document = PaperDocTemplate(str(PDF_OUT))
    story: list[Any] = [
        Paragraph(html.escape(title), styles["title"]),
        Paragraph(
            "Anonymous review manuscript | author details withheld",
            styles["subtitle"],
        ),
    ]
    cite_state: dict[str, int] = {}
    figure_count = 0
    table_count = 0
    index = 0
    while index < len(blocks):
        block = blocks[index]
        next_block = blocks[index + 1] if index + 1 < len(blocks) else None
        if block.kind == "heading":
            style = (
                styles["h1"]
                if block.level == 1
                else (styles["h2"] if block.level == 2 else styles["h3"])
            )
            story.append(Paragraph(pdf_inline(block.text, cite_state), style))
        elif block.kind == "paragraph":
            caption = block.text.startswith(("**Figure ", "**Table "))
            story.append(
                Paragraph(
                    pdf_inline(block.text, cite_state),
                    styles["caption"] if caption else styles["body"],
                )
            )
        elif block.kind in {"bullet", "number"}:
            bullet_text = "-" if block.kind == "bullet" else f"{block.ordinal}."
            story.append(
                Paragraph(
                    pdf_inline(block.text, cite_state),
                    styles["bullet"],
                    bulletText=bullet_text,
                )
            )
        elif block.kind == "equation":
            story.append(
                Paragraph(html.escape(display_equation(block.text)), styles["equation"])
            )
        elif block.kind == "image":
            image = make_pdf_image(block, document.width)
            figure_count += 1
            if next_block and next_block.kind == "paragraph" and next_block.text.startswith(
                "**Figure "
            ):
                caption = Paragraph(pdf_inline(next_block.text, cite_state), styles["caption"])
                story.append(KeepTogether([Spacer(1, 5), image, caption]))
                index += 1
            else:
                story.extend([Spacer(1, 5), image])
        elif block.kind == "table":
            table = make_pdf_table(block.rows, styles, document.width)
            table_count += 1
            if next_block and next_block.kind == "paragraph" and next_block.text.startswith(
                "**Table "
            ):
                caption = Paragraph(pdf_inline(next_block.text, cite_state), styles["caption"])
                story.append(KeepTogether([table, caption]))
                index += 1
            else:
                story.append(table)
        else:
            raise ValueError(f"Unsupported block: {block.kind}")
        index += 1

    entries = legacy.parse_bibtex(BIBLIOGRAPHY)
    story.append(Paragraph("References", styles["h1"]))
    for key, number in sorted(cite_state.items(), key=lambda item: item[1]):
        if key not in entries:
            raise KeyError(f"Missing bibliography entry: {key}")
        reference = html.escape(
            f"[{number}] {clean_reference(legacy.format_reference(entries[key]))}"
        )
        story.append(Paragraph(reference, styles["reference"]))
    document.build(story)
    reader = PdfReader(PDF_OUT)
    extracted = "\n".join(page.extract_text() or "" for page in reader.pages)
    normalized_extracted = " ".join(extracted.split())
    if title not in normalized_extracted or "References" not in normalized_extracted:
        raise AssertionError("PDF text audit did not recover the title and references heading")
    return {
        "path": str(PDF_OUT.relative_to(ROOT)),
        "sha256": sha256(PDF_OUT),
        "pages": len(reader.pages),
        "figures": figure_count,
        "tables": table_count,
        "citations": len(cite_state),
    }


def audit_docx(path: Path) -> dict[str, Any]:
    document = Document(path)
    drawings = sum(
        len(paragraph._p.xpath(".//w:drawing")) for paragraph in document.paragraphs
    )
    if drawings != EXPECTED_FIGURES or len(document.tables) != EXPECTED_TABLES:
        raise AssertionError(
            f"DOCX expected {EXPECTED_FIGURES} figures/{EXPECTED_TABLES} tables, "
            f"got {drawings}/{len(document.tables)}"
        )
    for table in document.tables:
        properties = table._tbl.tblPr
        width = properties.find(qn("w:tblW"))
        indent = properties.find(qn("w:tblInd"))
        if width is None or width.get(qn("w:type")) != "dxa":
            raise AssertionError("DOCX table lacks fixed DXA width")
        if indent is None or indent.get(qn("w:w")) != "120":
            raise AssertionError("DOCX table indent no longer matches the preset")
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    if "Abstract" not in text or "References" not in text:
        raise AssertionError("DOCX text audit did not recover Abstract and References")
    return {
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
        "inline_drawings": drawings,
    }


def build() -> dict[str, Any]:
    title, blocks = parse_markdown(MANUSCRIPT)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    docx = build_docx(title, blocks)
    pdf = build_pdf(title, blocks)
    docx_audit = audit_docx(DOCX_OUT)
    metadata = {
        "status": "pass",
        "source": str(MANUSCRIPT.relative_to(ROOT)),
        "source_sha256": sha256(MANUSCRIPT),
        "bibliography_sha256": sha256(BIBLIOGRAPHY),
        "design": {
            "base_preset": "narrative_proposal",
            "named_override": "journal-neutral scientific manuscript",
            "page": "US Letter portrait, 1 inch margins",
            "body": "Times New Roman 10.5 pt, justified, 1.2 spacing",
            "tables": "9360 DXA fixed geometry, 120 DXA indent",
            "images": "inline, maximum 6.35 inches",
        },
        "docx": docx,
        "docx_structural_audit": docx_audit,
        "pdf": pdf,
    }
    METADATA_OUT.write_text(json.dumps(metadata, indent=2, sort_keys=True), encoding="utf-8")
    return metadata


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check", action="store_true")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    if args.check:
        required = (DOCX_OUT, PDF_OUT, METADATA_OUT)
        missing = [str(path) for path in required if not path.is_file()]
        if missing:
            raise FileNotFoundError(f"Missing {ARTIFACT_LABEL} submission artifacts: {missing}")
        docx_audit = audit_docx(DOCX_OUT)
        reader = PdfReader(PDF_OUT)
        print(
            json.dumps(
                {
                    "status": "pass",
                    "docx": docx_audit,
                    "pdf_pages": len(reader.pages),
                },
                sort_keys=True,
            )
        )
        return
    print(json.dumps(build(), indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
